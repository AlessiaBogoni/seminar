import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Configuration
JS_FILE_PATH = "shared-data.js"
import os
from pptx import Presentation

THEME_PATH = "LEO_template.pptx"  # Verwende die neu abgespeicherte .pptx

prs = None
if os.path.exists(THEME_PATH):
    try:
        prs = Presentation(THEME_PATH)
        print(f"[SUCCESS] Vorlage '{THEME_PATH}' erfolgreich geladen!")
    except ValueError as e:
        print(f"[FEHLER] '{THEME_PATH}' konnte nicht als Präsentation gelesen werden: {e}")

if prs is None:
    print("[INFO] Erstelle eine neue leere Präsentation als Fallback.")
    prs = Presentation()

GROUP_NAMES = {
    "BS": "Bachelor Seminar Roadmap",
    "BT": "Bachelor Thesis Roadmap",
    "MS": "Master Seminar Roadmap"
}

BT_BLOCK_NAMES = {
    "1": "Block 1: Research Idea",
    "2": "Block 2: Research Question Finalization",
    "3": "Block 3: Experimental Design Setup",
    "4": "Block 4: Project Finalization",
    "5": "Block 5: Final Submission"
}

DEFAULT_BLOCK_NAMES = {
    "1": "Block 1: Orientation & Topic Assignment",
    "2": "Block 2: Research, Drafts & Writing Phase",
    "3": "Block 3: Final Submissions & Presentations"
}

BT_BLOCK_COLORS = {
    "1": "#8b5cf6", "2": "#5979a6", "3": "#10b981", "4": "#0696d4", "5": "#bc4899"
}

DEFAULT_BLOCK_COLORS = {
    "1": "#8b5cf6", "2": "#5979a6", "3": "#10b981"
}

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

# Load data
with open(JS_FILE_PATH, "r", encoding="utf-8") as f:
    js_content = f.read()

json_match = re.search(r"\[\s*\{.*\}\s*\]", js_content, re.DOTALL)
all_data = json.loads(json_match.group(0))

def parse_date(item):
    day, month, year = item["TargetDate"].split("-")
    return (int(year), int(month), int(day))

# Initialize presentation with the theme file
try:
    prs = Presentation(THEME_PATH)
except Exception:
    prs = Presentation()

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

CARDS_PER_SLIDE = 3

for group_code in ["BS", "BT", "MS"]:
    group_data = [item for item in all_data if item.get("Groups") == group_code]
    if not group_data:
        continue
        
    group_data.sort(key=parse_date)
    num_slides = (len(group_data) + CARDS_PER_SLIDE - 1) // CARDS_PER_SLIDE
    
    last_block = None
    
    for slide_idx in range(num_slides):
        chunk = group_data[slide_idx * CARDS_PER_SLIDE : (slide_idx + 1) * CARDS_PER_SLIDE]
        
        # Pull slide layout from loaded LEO theme
        layout_idx = 1 if len(prs.slide_layouts) > 1 else 0
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        
        # Remove body placeholder if present
        if len(slide.shapes.placeholders) > 1:
            content_shape = slide.shapes.placeholders[1]
            sp_elem = content_shape.element
            sp_elem.getparent().remove(sp_elem)
        
        # Slide Title from theme layout
        if slide.shapes.title:
            title_shape = slide.shapes.title
            title_text = GROUP_NAMES.get(group_code, f"{group_code} Roadmap")
            if num_slides > 1:
                title_text += f" ({slide_idx + 1}/{num_slides})"
            title_shape.text = title_text
        
        left_margin, top_margin = Inches(0.8), Inches(1.8)
        card_width, gap = Inches(11.733), Inches(0.25)
        n_items = len(chunk)
        card_height = (Inches(5.0) - (gap * (n_items - 1))) / n_items
        
        for idx, item in enumerate(chunk):
            raw_block = str(item.get("block", "1")).split('.')[0]
            hex_color = BT_BLOCK_COLORS.get(raw_block, "#8b5cf6") if group_code == "BT" else DEFAULT_BLOCK_COLORS.get(raw_block, "#8b5cf6")
            block_name = BT_BLOCK_NAMES.get(raw_block, f"Block {raw_block}") if group_code == "BT" else DEFAULT_BLOCK_NAMES.get(raw_block, f"Block {raw_block}")
                
            accent_rgb = hex_to_rgb(hex_color)
            card_y = top_margin + idx * (card_height + gap)
            
            # Card Container
            card_rect = slide.shapes.add_shape(1, left_margin, card_y, card_width, card_height)
            card_rect.fill.solid()
            card_rect.fill.fore_color.rgb = RGBColor(248, 250, 252)
            card_rect.line.color.rgb = RGBColor(226, 232, 240)
            
            # Left Vertical Accent Stripe
            accent_stripe = slide.shapes.add_shape(1, left_margin, card_y, Inches(0.12), card_height)
            accent_stripe.fill.solid()
            accent_stripe.fill.fore_color.rgb = accent_rgb
            accent_stripe.line.fill.background()
            
            # Text Frame
            tb = slide.shapes.add_textbox(left_margin + Inches(0.3), card_y + Inches(0.12), card_width - Inches(0.5), card_height - Inches(0.24))
            tf = tb.text_frame
            tf.word_wrap = True
            
            # Block Header Title on Switch
            if raw_block != last_block:
                p_block = tf.paragraphs[0]
                p_block.text = block_name.upper()
                p_block.font.bold = True
                p_block.font.size = Pt(13)
                p_block.font.color.rgb = accent_rgb
                p_block.space_after = Pt(4)
                last_block = raw_block
                p_head = tf.add_paragraph()
            else:
                p_head = tf.paragraphs[0]
            
            # Date & Title (20pt)
            p_head.text = f"{item['TargetDate']}   |   {item['MilestoneName'].strip()}"
            p_head.font.bold = True
            p_head.font.size = Pt(20)
            p_head.font.color.rgb = RGBColor(15, 23, 42)
            p_head.space_after = Pt(4)
            
            # Description (16pt)
            if "description" in item and item["description"]:
                p_desc = tf.add_paragraph()
                p_desc.text = item["description"]
                p_desc.font.size = Pt(16)
                p_desc.font.color.rgb = RGBColor(51, 65, 85)

prs.save("roadmaps_leo_theme.pptx")